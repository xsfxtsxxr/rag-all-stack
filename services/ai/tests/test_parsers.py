import base64
from io import BytesIO

import pytest
from docx import Document

from app.parsers import registry

# 1x1 透明 PNG，供纯图片 docx 测试样本使用
_TINY_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJ"
    "AAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg=="
)


def test_registry_supports_docx() -> None:
    assert registry.get("合同.docx").supports("合同.docx")


def test_registry_rejects_unknown() -> None:
    with pytest.raises(ValueError):
        registry.get("笔记.txt")


def test_docx_parse_paragraphs_and_tables() -> None:
    doc = Document()
    doc.add_paragraph("第一条：产品规格")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "型号"
    table.rows[0].cells[1].text = "VFD-220"
    buf = BytesIO()
    doc.save(buf)

    text = registry.get("报价单.docx").parse(buf.getvalue(), "报价单.docx")
    assert "产品规格" in text
    assert "VFD-220" in text


def test_docx_image_only_falls_back_to_ocr(monkeypatch: pytest.MonkeyPatch) -> None:
    doc = Document()
    doc.add_picture(BytesIO(_TINY_PNG))
    buf = BytesIO()
    doc.save(buf)

    monkeypatch.setattr("app.ocr.ocr_image", lambda _content: "红外窗口 磁吸式")
    text = registry.get("红外窗口.docx").parse(buf.getvalue(), "红外窗口.docx")
    assert "红外窗口" in text
