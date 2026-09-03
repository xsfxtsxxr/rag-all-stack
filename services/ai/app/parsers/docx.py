from io import BytesIO
from zipfile import ZipFile

from docx import Document

from app.parsers.base import DocumentParser

_MEDIA_PREFIX = "word/media/"
_RASTER_IMAGE_EXTENSIONS = (".jpg", ".jpeg", ".png", ".webp", ".bmp")


class DocxParser(DocumentParser):
    """Word 文档解析（python-docx）：段落 + 表格单元格文本；纯图片文档回退 OCR"""

    def supports(self, filename: str) -> bool:
        return filename.lower().endswith(".docx")

    def parse(self, content: bytes, filename: str) -> str:
        del filename  # 保留接口一致性
        doc = Document(BytesIO(content))
        parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
        text = "\n".join(parts)
        return text if text else self._ocr_embedded_images(content)

    @staticmethod
    def _ocr_embedded_images(content: bytes) -> str:
        """版式内无文字（纯图片页）时，逐张 OCR docx 内嵌位图并拼接文本"""
        from app.ocr import ocr_image  # 延迟导入：避免启动即加载 OCR 模型依赖

        with ZipFile(BytesIO(content)) as archive:
            images = [
                archive.read(name)
                for name in archive.namelist()
                if name.startswith(_MEDIA_PREFIX)
                and name.lower().endswith(_RASTER_IMAGE_EXTENSIONS)
            ]
        pages: list[str] = []
        for image in images:
            try:
                page_text = ocr_image(image)
            except ValueError:
                continue  # 单张图片解码失败不影响其余页面
            if page_text.strip():
                pages.append(page_text)
        return "\n".join(pages)
